from openpyxl.worksheet.worksheet import Worksheet


def parse_down(sheet: Worksheet, i, j, r, c, ftype):
    course_map = {}
    if ftype == 1:
        while (i <= r):
            v = sheet.cell(i, j).value
            v2 = sheet.cell(i, j+1).value
            s = sheet.cell(i, j+2).value
            i+=1

            if (v2 == "(19M21HS111)"):
                course_map.update({str(v2).strip("() "): str(s)})

            if (v is None or v2 is None or s is None):
                continue
            
            v = str(v).strip("() ")
            v2 = str(v2).strip("() ")
            s = str(s).strip()
            course_map.update({v: s, v2: s})
    elif ftype == 2:
        while i <= r:
            v = str(sheet.cell(i, j).value)
            s = str(sheet.cell(i, j+1).value)
            if v == "Faculty Abbreviation with Names":
                break

            i+=1
            if (v is None or s is None):
                continue

            v1, v2 = v.split("/")
            course_map.update({v1.strip(): s, v2.strip(): s})
    return course_map

def parse_courses(sheet: Worksheet, row: int, col: int):
    course_map = {}
    for i in range(1, row+1):
        for j in range(1, col+1):
            value = str(sheet.cell(i, j).value)
            nvalue = str(sheet.cell(i, j+1).value)
            if value in ["Short Subject Code", "CODE", "SHORT FORM"] and (nvalue in ["Subject Code", "SUBJECT CODE"]):
                course_map.update(parse_down(sheet, i+1, j, row, col, ftype=1))
            elif value in ["SHORT FORM / SUBJECT CODE"] and nvalue in ["SUBJECT NAME"]:
                course_map.update(parse_down(sheet, i+1, j, row, col, ftype=2))
            
    return course_map

